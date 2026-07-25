"""Build the Excel workbook with live formulas (openpyxl)."""
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

tally = pd.read_csv("/home/claude/wc/pipe_tally.csv")
srv = pd.read_csv("/home/claude/wc/directional_survey.csv")
S = json.load(open("/home/claude/wc/summary.json"))

ARIAL = "Arial"
BLUE = Font(name=ARIAL, size=10, color="0000FF")          # hardcoded input
BLACK = Font(name=ARIAL, size=10)                          # formula
BOLD = Font(name=ARIAL, size=10, bold=True)
TITLE = Font(name=ARIAL, size=14, bold=True)
SECTION = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")
NOTE = Font(name=ARIAL, size=9, italic=True, color="595959")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="2E5C8A")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()

# ============================================================ Well_Data
ws = wb.active
ws.title = "Well_Data"
R = {}   # label -> row number

ws["A1"] = "WELL CONTROL TEST DATA — Test-1H (synthetic)"
ws["A1"].font = TITLE
ws["A2"] = ("Synthetic dataset for exercising a kill planner. Not from a real well. "
            "Blue = hardcoded input · Black = formula.")
ws["A2"].font = NOTE

hdr = ["Parameter", "Value", "Units", "Source / Note"]
for c, h in enumerate(hdr, start=1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
    cell.fill = HDR_FILL
    cell.border = BOX

row = 5


def section(title):
    global row
    ws.cell(row=row, column=1, value=title).font = SECTION
    for c in range(1, 5):
        ws.cell(row=row, column=c).fill = SEC_FILL
    row += 1


def put(label, value, units="", note="", fmt=None, formula=False):
    global row
    ws.cell(row=row, column=1, value=label).font = BLACK
    c = ws.cell(row=row, column=2, value=value)
    c.font = BLACK if formula else BLUE
    if fmt:
        c.number_format = fmt
    ws.cell(row=row, column=3, value=units).font = BLACK
    ws.cell(row=row, column=4, value=note).font = NOTE
    for cc in range(1, 5):
        ws.cell(row=row, column=cc).border = BOX
    R[label] = row
    row += 1


def ref(label):
    return f"'Well_Data'!$B${R[label]}"


n_srv = len(srv)
SRV_FIRST, SRV_LAST = 2, n_srv + 1          # data rows on Directional_Survey
n_tal = len(tally)
TAL_FIRST, TAL_LAST = 2, n_tal + 1
VOL_FIRST, VOL_LAST = 2, n_tal + 1

section("WELL AND RIG")
put("Well Name", "Test-1H (synthetic)", "", "Fictitious. Generated test data.")
put("Rig", "Land rig, triples", "", "")
put("Unit System", "Field (ft, in, ppg, psi, bbl)", "", "")

section("HOLE GEOMETRY")
put("Bit Depth (MD)", S["Bit_Depth_MD_ft"], "ft",
    "Equals total tally length.", "0.00")
put("Last Survey (MD)", f"=MAX('Directional_Survey'!$A${SRV_FIRST}:$A${SRV_LAST})",
    "ft", "Deepest survey station.", "0.00", formula=True)
put("Last Survey (TVD)",
    f"=INDEX('Directional_Survey'!$D${SRV_FIRST}:$D${SRV_LAST},"
    f"MATCH({ref('Last Survey (MD)')},'Directional_Survey'!$A${SRV_FIRST}:$A${SRV_LAST},0))",
    "ft", "From survey sheet.", "0.00", formula=True)
put("Last Survey Inclination",
    f"=INDEX('Directional_Survey'!$B${SRV_FIRST}:$B${SRV_LAST},"
    f"MATCH({ref('Last Survey (MD)')},'Directional_Survey'!$A${SRV_FIRST}:$A${SRV_LAST},0))",
    "deg", "", "0.00", formula=True)
put("Bit Depth (TVD)",
    f"={ref('Last Survey (TVD)')}+({ref('Bit Depth (MD)')}-{ref('Last Survey (MD)')})"
    f"*COS(RADIANS({ref('Last Survey Inclination')}))",
    "ft", "Extended below last station on last inclination.", "0.00", formula=True)
put("Kick-off Point (KOP)", S["KOP_ft_MD"], "ft MD", "", "0.00")
put("Build Rate", S["Build_Rate_deg_per_100ft"], "deg/100ft", "", "0.00")
put("Hold Inclination", S["Hold_Inclination_deg"], "deg", "", "0.00")
put("End of Build (EOB)",
    f"={ref('Kick-off Point (KOP)')}+{ref('Hold Inclination')}"
    f"/{ref('Build Rate')}*100", "ft MD", "", "0.00", formula=True)

section("CASING AND OPEN HOLE")
put("Surface Casing", S["Surface_Casing"], "", "")
put("Surface Shoe (MD)", S["Surface_Shoe_MD_ft"], "ft", "", "0.00")
put("Intermediate Casing", S["Intermediate_Casing"], "",
    "Governs MAASP — this is the exposed shoe.")
put("Intermediate Shoe (MD)", S["Intermediate_Shoe_MD_ft"], "ft",
    "A survey station is taken at this depth.", "0.00")
put("Intermediate Shoe (TVD)",
    f"=INDEX('Directional_Survey'!$D${SRV_FIRST}:$D${SRV_LAST},"
    f"MATCH({ref('Intermediate Shoe (MD)')},"
    f"'Directional_Survey'!$A${SRV_FIRST}:$A${SRV_LAST},0))",
    "ft", "Exact station lookup, no interpolation.", "0.00", formula=True)
put("Casing Drift ID", S["Casing_ID_in"], "in", "9-5/8in 47.0 ppf N-80.", "0.000")
put("Open Hole Size", S["Open_Hole_Size_in"], "in", "", "0.000")

section("MUD AND PUMP")
put("Mud Weight (current)", S["Mud_Weight_ppg"], "ppg", "", "0.00")
put("Formation Integrity (LOT)", S["LOT_EMW_ppg"], "ppg EMW",
    "At intermediate shoe.", "0.00")
put("Pump", S["Pump_Type"], "", "")
put("Pump Liner", 6.5, "in", "", "0.00")
put("Pump Stroke Length", 12.0, "in", "", "0.00")
put("Pump Volumetric Efficiency", 0.95, "", "", "0.0%")
put("Pump Output",
    f"=0.000243*{ref('Pump Liner')}^2*{ref('Pump Stroke Length')}"
    f"*{ref('Pump Volumetric Efficiency')}",
    "bbl/stk", "Triplex: 0.000243 x liner^2 x stroke x eff.", "0.00000", formula=True)
put("Kill Rate", S["Kill_Rate_spm"], "spm", "", "0")
put("Slow Circulating Pressure", S["Slow_Circulating_Pressure_psi"], "psi",
    "Measured at kill rate.", "0")
put("Surface Line Volume", S["Surface_Line_Volume_bbl"], "bbl",
    "Standpipe + hose + kelly. Excluded from strokes-to-bit.", "0.00")

section("VOLUMES AND STROKES")
put("String Internal Volume",
    f"=SUM('Volumes'!$M${VOL_FIRST}:$M${VOL_LAST})", "bbl", "", "0.00", formula=True)
put("Annulus Volume — Open Hole",
    f"=SUM('Volumes'!$H${VOL_FIRST}:$H${VOL_LAST})", "bbl",
    "Shoe to bit.", "0.00", formula=True)
put("Annulus Volume — Inside Casing",
    f"=SUM('Volumes'!$J${VOL_FIRST}:$J${VOL_LAST})", "bbl",
    "Surface to shoe.", "0.00", formula=True)
put("Total Annulus Volume",
    f"={ref('Annulus Volume — Open Hole')}+{ref('Annulus Volume — Inside Casing')}",
    "bbl", "", "0.00", formula=True)
put("Total Hole Volume",
    f"={ref('String Internal Volume')}+{ref('Total Annulus Volume')}",
    "bbl", "", "0.00", formula=True)
put("Strokes — Surface to Bit",
    f"={ref('String Internal Volume')}/{ref('Pump Output')}", "stk", "", "0", formula=True)
put("Strokes — Bit to Shoe",
    f"={ref('Annulus Volume — Open Hole')}/{ref('Pump Output')}", "stk", "", "0", formula=True)
put("Strokes — Bit to Surface",
    f"={ref('Total Annulus Volume')}/{ref('Pump Output')}", "stk", "", "0", formula=True)
put("Strokes — Full Circulation",
    f"={ref('Strokes — Surface to Bit')}+{ref('Strokes — Bit to Surface')}",
    "stk", "", "0", formula=True)
put("Time — Surface to Bit",
    f"={ref('Strokes — Surface to Bit')}/{ref('Kill Rate')}", "min", "", "0.0", formula=True)
put("Time — Full Circulation",
    f"={ref('Strokes — Full Circulation')}/{ref('Kill Rate')}", "min", "", "0.0", formula=True)

section("SHUT-IN DATA (the scenario to solve)")
put("SIDPP", S["SIDPP_psi"], "psi", "Shut-in drillpipe pressure, stabilised.", "0")
put("SICP", S["SICP_psi"], "psi",
    "Shut-in casing pressure. Back-calculated from the influx below so the "
    "dataset is self-consistent.", "0")
put("Pit Gain", S["Pit_Gain_bbl"], "bbl", "", "0.0")
put("Influx Gradient (assumed)", S["Influx_Gradient_psi_per_ft"], "psi/ft",
    "Gas. Used only to derive SICP.", "0.000")
put("Influx Height (MD)", S["Influx_Height_MD_ft"], "ft",
    "Pit gain filled up the annulus from the bit.", "0.0")
put("Influx Height (TVD)", S["Influx_Height_TVD_ft"], "ft", "", "0.0")
put("Influx Top (MD)", S["Influx_Top_MD_ft"], "ft", "", "0.0")

section("EXPECTED KILL SHEET RESULTS")
put("Hydrostatic at Bit (TVD)",
    f"=0.052*{ref('Mud Weight (current)')}*{ref('Bit Depth (TVD)')}",
    "psi", "", "0", formula=True)
put("Bottomhole Pressure",
    f"={ref('Hydrostatic at Bit (TVD)')}+{ref('SIDPP')}", "psi", "", "0", formula=True)
put("Kill Mud Weight",
    f"={ref('Mud Weight (current)')}+{ref('SIDPP')}"
    f"/(0.052*{ref('Bit Depth (TVD)')})", "ppg", "", "0.00", formula=True)
put("Initial Circulating Pressure (ICP)",
    f"={ref('Slow Circulating Pressure')}+{ref('SIDPP')}", "psi", "", "0", formula=True)
put("Final Circulating Pressure (FCP)",
    f"={ref('Slow Circulating Pressure')}*{ref('Kill Mud Weight')}"
    f"/{ref('Mud Weight (current)')}", "psi", "", "0", formula=True)
put("MAASP — current mud weight",
    f"=({ref('Formation Integrity (LOT)')}-{ref('Mud Weight (current)')})"
    f"*0.052*{ref('Intermediate Shoe (TVD)')}", "psi", "", "0", formula=True)
put("MAASP — kill mud weight",
    f"=({ref('Formation Integrity (LOT)')}-{ref('Kill Mud Weight')})"
    f"*0.052*{ref('Intermediate Shoe (TVD)')}", "psi", "", "0", formula=True)
put("SICP vs MAASP margin",
    f"={ref('MAASP — current mud weight')}-{ref('SICP')}", "psi",
    "Positive means the shut-in pressure is below the shoe limit.", "0", formula=True)

ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 11
ws.column_dimensions["D"].width = 62
ws.freeze_panes = "A5"


# ============================================================ helper
def write_table(ws_, df, formats=None, widths=None):
    for c, col in enumerate(df.columns, start=1):
        cell = ws_.cell(row=1, column=c, value=col)
        cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BOX
    for r in range(len(df)):
        for c, col in enumerate(df.columns, start=1):
            v = df.iloc[r][col]
            if pd.isna(v):
                v = None
            cell = ws_.cell(row=r + 2, column=c, value=v)
            cell.font = BLUE
            if formats and col in formats:
                cell.number_format = formats[col]
    if widths:
        for col, w in widths.items():
            ws_.column_dimensions[get_column_letter(df.columns.get_loc(col) + 1)].width = w
    ws_.freeze_panes = "A2"


# ============================================================ Pipe_Tally
wt = wb.create_sheet("Pipe_Tally")
tal_cols = list(tally.columns)
write_table(wt, tally[["Item_No", "Joint_No", "Stand_No", "Description",
                       "Component_Type", "OD_in", "ID_in", "Nominal_Weight_ppf",
                       "Connection", "Length_ft"]].assign(
    Cum_Length_ft=None, Top_Depth_ft=None, Bottom_Depth_ft=None,
    Capacity_bbl_per_ft=None, Capacity_bbl=None,
    Displacement_bbl_per_ft=None, Displacement_bbl=None)[tal_cols],
    formats={"OD_in": "0.000", "ID_in": "0.000", "Nominal_Weight_ppf": "0.00",
             "Length_ft": "0.00"},
    widths={"Description": 38, "Component_Type": 14, "Connection": 18})

td = ref("Bit Depth (MD)")
for i in range(n_tal):
    r = i + 2
    p = r - 1
    wt[f"K{r}"] = f"=J{r}" if i == 0 else f"=K{p}+J{r}"
    wt[f"L{r}"] = f"={td}-K{r}"
    wt[f"M{r}"] = f"={td}-K{r}+J{r}"
    wt[f"N{r}"] = f"=G{r}^2/1029.4"
    wt[f"O{r}"] = f"=N{r}*J{r}"
    wt[f"P{r}"] = f"=(F{r}^2-G{r}^2)/1029.4"
    wt[f"Q{r}"] = f"=P{r}*J{r}"
    for col, fmt in (("K", "0.00"), ("L", "0.00"), ("M", "0.00"),
                     ("N", "0.00000"), ("O", "0.000"),
                     ("P", "0.00000"), ("Q", "0.000")):
        wt[f"{col}{r}"].font = BLACK
        wt[f"{col}{r}"].number_format = fmt

tot = n_tal + 2
wt[f"D{tot}"] = "TOTAL"
wt[f"D{tot}"].font = BOLD
wt[f"J{tot}"] = f"=SUM(J{TAL_FIRST}:J{TAL_LAST})"
wt[f"O{tot}"] = f"=SUM(O{TAL_FIRST}:O{TAL_LAST})"
wt[f"Q{tot}"] = f"=SUM(Q{TAL_FIRST}:Q{TAL_LAST})"
for col, fmt in (("J", "0.00"), ("O", "0.00"), ("Q", "0.00")):
    wt[f"{col}{tot}"].font = BOLD
    wt[f"{col}{tot}"].number_format = fmt
wt.cell(row=tot + 2, column=4,
        value="Displacement is plain pipe body only; it excludes tool joint steel, "
              "so real closed-end displacement runs slightly higher.").font = NOTE

# ============================================================ Directional_Survey
wd = wb.create_sheet("Directional_Survey")
srv_cols = list(srv.columns)
write_table(wd, srv[["MD_ft", "Inclination_deg", "Azimuth_deg"]].assign(
    TVD_ft=None, Northing_ft=None, Easting_ft=None, DLS_deg_per_100ft=None,
    Course_Length_ft=None, Vertical_Section_ft=None, Closure_Distance_ft=None,
    Closure_Azimuth_deg=None)[srv_cols],
    formats={"MD_ft": "0.00", "Inclination_deg": "0.00", "Azimuth_deg": "0.00"},
    widths={"MD_ft": 11, "Inclination_deg": 13, "Azimuth_deg": 12})

section("SURVEY REFERENCE")
ws.cell(row=row, column=1, value="Vertical Section Azimuth").font = BLACK
c = ws.cell(row=row, column=2, value=144.0)
c.font = BLUE
c.number_format = "0.00"
ws.cell(row=row, column=3, value="deg").font = BLACK
ws.cell(row=row, column=4,
        value="Reference azimuth for the vertical section column.").font = NOTE
R["Vertical Section Azimuth"] = row
vs_az = ref("Vertical Section Azimuth")
row += 1

for i in range(len(srv)):
    r = i + 2
    p = r - 1
    if i == 0:
        wd[f"D{r}"], wd[f"E{r}"], wd[f"F{r}"] = 0, 0, 0
        wd[f"G{r}"], wd[f"H{r}"] = 0, 0
    else:
        dl = (f"ACOS(MIN(1,MAX(-1,COS(RADIANS(B{r}-B{p}))"
              f"-SIN(RADIANS(B{p}))*SIN(RADIANS(B{r}))*(1-COS(RADIANS(C{r}-C{p}))))))")
        rf = f"IF({dl}=0,1,2/({dl})*TAN(({dl})/2))"
        wd[f"H{r}"] = f"=A{r}-A{p}"
        wd[f"D{r}"] = (f"=D{p}+(A{r}-A{p})/2*(COS(RADIANS(B{p}))+COS(RADIANS(B{r})))*({rf})")
        wd[f"E{r}"] = (f"=E{p}+(A{r}-A{p})/2*(SIN(RADIANS(B{p}))*COS(RADIANS(C{p}))"
                       f"+SIN(RADIANS(B{r}))*COS(RADIANS(C{r})))*({rf})")
        wd[f"F{r}"] = (f"=F{p}+(A{r}-A{p})/2*(SIN(RADIANS(B{p}))*SIN(RADIANS(C{p}))"
                       f"+SIN(RADIANS(B{r}))*SIN(RADIANS(C{r})))*({rf})")
        wd[f"G{r}"] = f"=IF(H{r}=0,0,DEGREES({dl})*100/H{r})"
    wd[f"J{r}"] = f"=SQRT(E{r}^2+F{r}^2)"
    wd[f"K{r}"] = f"=IF(J{r}=0,0,MOD(DEGREES(ATAN2(E{r},F{r})),360))"
    wd[f"I{r}"] = f"=J{r}*COS(RADIANS(K{r}-{vs_az}))"
    for col, fmt in (("D", "0.00"), ("E", "0.00"), ("F", "0.00"), ("G", "0.00"),
                     ("H", "0.00"), ("I", "0.00"), ("J", "0.00"), ("K", "0.00")):
        wd[f"{col}{r}"].font = BLACK
        wd[f"{col}{r}"].number_format = fmt

wd.cell(row=len(srv) + 3, column=1,
        value="TVD / Northing / Easting / DLS computed by minimum curvature from the "
              "MD, inclination and azimuth columns.").font = NOTE

# ============================================================ Volumes
wv = wb.create_sheet("Volumes")
vhdr = ["Item_No", "Description", "OD_in", "ID_in", "Top_Depth_ft", "Bottom_Depth_ft",
        "Ann_Len_OpenHole_ft", "Ann_Vol_OpenHole_bbl", "Ann_Len_InCasing_ft",
        "Ann_Vol_InCasing_bbl", "Internal_Cap_bbl_per_ft", "Internal_Len_ft",
        "Internal_Vol_bbl"]
for c, h in enumerate(vhdr, start=1):
    cell = wv.cell(row=1, column=c, value=h)
    cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
    cell.fill = HDR_FILL
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
    cell.border = BOX

shoe = ref("Intermediate Shoe (MD)")
ohd = ref("Open Hole Size")
cid = ref("Casing Drift ID")
for i in range(n_tal):
    r = i + 2
    tr = i + 2   # same row on Pipe_Tally
    wv[f"A{r}"] = f"='Pipe_Tally'!A{tr}"
    wv[f"B{r}"] = f"='Pipe_Tally'!D{tr}"
    wv[f"C{r}"] = f"='Pipe_Tally'!F{tr}"
    wv[f"D{r}"] = f"='Pipe_Tally'!G{tr}"
    wv[f"E{r}"] = f"='Pipe_Tally'!L{tr}"
    wv[f"F{r}"] = f"='Pipe_Tally'!M{tr}"
    wv[f"G{r}"] = f"=MAX(0,MIN(F{r},{td})-MAX(E{r},{shoe}))"
    wv[f"H{r}"] = f"=G{r}*({ohd}^2-C{r}^2)/1029.4"
    wv[f"I{r}"] = f"=MAX(0,MIN(F{r},{shoe})-E{r})"
    wv[f"J{r}"] = f"=I{r}*({cid}^2-C{r}^2)/1029.4"
    wv[f"K{r}"] = f"=D{r}^2/1029.4"
    wv[f"L{r}"] = f"=F{r}-E{r}"
    wv[f"M{r}"] = f"=K{r}*L{r}"
    for col, fmt in (("C", "0.000"), ("D", "0.000"), ("E", "0.00"), ("F", "0.00"),
                     ("G", "0.00"), ("H", "0.000"), ("I", "0.00"), ("J", "0.000"),
                     ("K", "0.00000"), ("L", "0.00"), ("M", "0.000")):
        wv[f"{col}{r}"].font = BLACK
        wv[f"{col}{r}"].number_format = fmt
    wv[f"A{r}"].font = BLACK
    wv[f"B{r}"].font = BLACK

vt = n_tal + 2
wv[f"B{vt}"] = "TOTAL"
wv[f"B{vt}"].font = BOLD
for col in ("H", "J", "M"):
    wv[f"{col}{vt}"] = f"=SUM({col}{VOL_FIRST}:{col}{VOL_LAST})"
    wv[f"{col}{vt}"].font = BOLD
    wv[f"{col}{vt}"].number_format = "0.00"
wv.column_dimensions["B"].width = 38
for col in ("G", "H", "I", "J", "K", "L", "M"):
    wv.column_dimensions[col].width = 13

# ============================================================ Notes
wn = wb.create_sheet("Notes")
wn["A1"] = "NOTES, ASSUMPTIONS AND LEGEND"
wn["A1"].font = TITLE
notes = [
    "",
    "PURPOSE",
    "Synthetic test data for a well control kill planner. Nothing here came from a real "
    "well; the numbers are internally consistent so a correct planner should reproduce the "
    "expected results on the Well_Data sheet.",
    "",
    "LEGEND",
    "Blue text = hardcoded input value. Black text = formula. Every derived figure "
    "recalculates from the inputs, so you can change a value and watch it propagate.",
    "",
    "WELL",
    "Build-and-hold land well. Vertical to 2,000 ft MD, build at 2.5 deg/100 ft to 35 deg, "
    "hold to 12,500 ft MD. 13-3/8in surface shoe at 1,800 ft; 9-5/8in intermediate shoe at "
    "8,200 ft; 8-1/2in open hole below. A survey station falls exactly on the intermediate "
    "shoe so shoe TVD needs no interpolation.",
    "",
    "SURVEY",
    "Minimum curvature. Stations every 200 ft while vertical, every 30 ft through the build, "
    "every 93 ft (one stand) in the tangent. Last station 12,460 ft MD; the bit is at 12,500 "
    "ft, so TVD at the bit is extended below the last station on the last inclination — the "
    "same thing a rig would do. Inclination and azimuth carry small station-to-station "
    "variation, so the tangent section is not a perfect 35.00 deg.",
    "",
    "TALLY",
    "405 items, bit to surface, totalling exactly 12,500.00 ft. Individual joint lengths vary "
    "as measured lengths do. A 5.07 ft pup joint at the top makes the tally close on depth. "
    "Capacity = ID^2/1029.4. Displacement = (OD^2-ID^2)/1029.4, which is plain pipe body "
    "only and excludes tool joint steel — real closed-end displacement is slightly higher.",
    "",
    "KICK SCENARIO",
    "22 bbl pit gain, filled up the annulus from the bit, giving a 698.6 ft (MD) influx "
    "column. SIDPP of 480 psi is the primary input. SICP was then back-calculated assuming a "
    "0.12 psi/ft gas influx, so SIDPP, SICP and pit gain agree with each other rather than "
    "being three unrelated numbers. Shut-in casing pressure sits comfortably below MAASP, so "
    "the well is safely shut in and the kill is feasible.",
    "",
    "FORMULAS USED",
    "Kill mud weight = MW + SIDPP / (0.052 x TVD)",
    "ICP = slow circulating pressure + SIDPP",
    "FCP = slow circulating pressure x kill MW / current MW",
    "MAASP = (LOT EMW - MW) x 0.052 x shoe TVD",
    "Pump output (triplex) = 0.000243 x liner^2 x stroke x efficiency",
    "Annular capacity = (hole ID^2 - pipe OD^2) / 1029.4",
    "",
    "WHAT TO CHECK IN YOUR PLANNER",
    "TVD at the bit, TVD at the shoe, kill mud weight, ICP, FCP, MAASP, strokes from surface "
    "to bit, strokes from bit to shoe, and total circulation strokes. All are on the "
    "Well_Data sheet under EXPECTED KILL SHEET RESULTS.",
]
r = 3
for line in notes:
    c = wn.cell(row=r, column=1, value=line)
    c.font = BOLD if (line.isupper() and line) else Font(name=ARIAL, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
wn.column_dimensions["A"].width = 118

wb.save("/home/claude/wc/well_control_test_data.xlsx")
print("written")
print("Well_Data rows:", {k: v for k, v in list(R.items())[:3]}, "... total", len(R))
